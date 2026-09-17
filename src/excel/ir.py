from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .cells import (
    cell_str,
    filled_count,
    is_empty_row,
    looks_like_header,
    mostly_numeric,
    row_width,
    trim_rows,
)

SheetKind = str


@dataclass
class TableBlock:
    sheet: str
    headers: list[str]
    rows: list[tuple[int, list[Any]]]
    start_row: int
    end_row: int
    kind: SheetKind = "data_table"
    column_map: dict[str, dict[str, str]] = field(default_factory=dict)
    header_rows: list[list[Any]] = field(default_factory=list)


def load_workbook_tables(path: Path) -> list[TableBlock]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl required for xlsx parse — pip install openpyxl") from exc

    wb = load_workbook(path, data_only=True)
    tables: list[TableBlock] = []
    try:
        for sheet in wb.worksheets:
            grid = sheet_grid(sheet)
            tables.extend(tables_from_grid(sheet.title, grid))
    finally:
        wb.close()
    if not tables:
        raise RuntimeError(f"no data tables found in {path}")
    return tables


def sheet_grid(sheet: Any) -> list[list[Any]]:
    max_row = sheet.max_row or 0
    max_col = sheet.max_column or 0
    if max_row < 1 or max_col < 1:
        return []
    grid = [[None] * max_col for _ in range(max_row)]
    for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            grid[cell.row - 1][cell.column - 1] = cell.value
    for merged in sheet.merged_cells.ranges:
        value = grid[merged.min_row - 1][merged.min_col - 1]
        for row_i in range(merged.min_row, merged.max_row + 1):
            for col_i in range(merged.min_col, merged.max_col + 1):
                grid[row_i - 1][col_i - 1] = value
    return grid


def tables_from_grid(sheet_name: str, grid: list[list[Any]]) -> list[TableBlock]:
    if not grid:
        return []
    tables: list[TableBlock] = []
    i = 0
    while i < len(grid):
        while i < len(grid) and is_empty_row(grid[i]):
            i += 1
        if i >= len(grid):
            break
        header_start = skip_title_rows(grid, i)
        header_end = header_band_end(grid, header_start)
        raw_headers = [list(row) for row in grid[header_start:header_end]]
        headers = join_headers(raw_headers)
        data_start = header_end
        data_end = data_start
        while data_end < len(grid) and not is_empty_row(grid[data_end]):
            data_end += 1
        rows = [
            (idx + 1, list(grid[idx]))
            for idx in range(data_start, data_end)
            if not is_empty_row(grid[idx])
        ]
        width = max(
            len(headers),
            row_width(rows),
            max((len(row) for row in raw_headers), default=1),
        )
        if headers or rows:
            tables.append(
                TableBlock(
                    sheet=sheet_name,
                    headers=headers or [f"col_{n}" for n in range(row_width(rows))],
                    rows=trim_rows(rows, width),
                    start_row=header_start + 1,
                    end_row=data_end,
                    header_rows=raw_headers,
                )
            )
        i = data_end + 1
    return tables


def skip_title_rows(grid: list[list[Any]], start: int) -> int:
    i = start
    while i < len(grid) - 1:
        filled = filled_count(grid[i])
        nxt = filled_count(grid[i + 1])
        if filled <= 1 and nxt >= max(2, filled + 1) and not mostly_numeric(grid[i + 1]):
            i += 1
            continue
        break
    return i


def header_band_end(grid: list[list[Any]], start: int) -> int:
    end = start + 1
    while end < len(grid) and looks_like_header(grid[end]) and not is_empty_row(grid[end]):
        if mostly_numeric(grid[end]):
            break
        end += 1
        if end - start >= 4:
            break
    return max(end, start + 1)


def join_headers(header_rows: list[list[Any]]) -> list[str]:
    if not header_rows:
        return []
    width = max(len(row) for row in header_rows)
    headers: list[str] = []
    for col in range(width):
        parts: list[str] = []
        for row in header_rows:
            value = cell_str(row[col] if col < len(row) else "")
            if value and value not in parts:
                parts.append(value)
        headers.append(" / ".join(parts) if parts else f"col_{col}")
    while len(headers) > 1 and headers[-1].startswith("col_"):
        headers.pop()
    return headers
