from __future__ import annotations

import gzip
import json
import shutil
from pathlib import Path
from typing import Any

from ub_local.pipeline.types import FlatChunk


def parse_survey_input(
    input_path: Path | str,
    output_dir: Path | str,
    *,
    log: Any = print,
) -> Path:
    """
    Normalize survey inputs into chunks.jsonl.gz for embed_seed_bundle.

    Accepts:
    - *.jsonl.gz / *.jsonl  (already flat chunks)
    - *.xlsx / *.xlsm       (one chunk per data row)
    - *.docx                (one chunk per non-empty paragraph)
    """
    src = Path(input_path).expanduser().resolve()
    if not src.is_file():
        raise FileNotFoundError(src)
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    dest = out / "chunks.jsonl.gz"
    lower = src.name.lower()

    if lower.endswith(".jsonl.gz") or lower.endswith(".chunks.jsonl.gz"):
        shutil.copy2(src, dest)
        log(f"copied prebuilt chunks → {dest}")
        return dest

    if lower.endswith(".jsonl"):
        chunks = _load_jsonl(src)
        _write_jsonl_gz(dest, chunks)
        log(f"packed {len(chunks)} jsonl rows → {dest}")
        return dest

    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        chunks = _xlsx_to_chunks(src)
        _write_jsonl_gz(dest, chunks)
        log(f"xlsx → {len(chunks)} chunks → {dest}")
        return dest

    if lower.endswith(".docx"):
        chunks = _docx_to_chunks(src)
        _write_jsonl_gz(dest, chunks)
        log(f"docx → {len(chunks)} chunks → {dest}")
        return dest

    raise RuntimeError(
        "unsupported survey input; use .xlsx, .docx, .jsonl, or .jsonl.gz "
        f"(got {src.suffix})"
    )


def _load_jsonl(path: Path) -> list[FlatChunk]:
    rows: list[FlatChunk] = []
    with path.open("rt", encoding="utf-8") as f:
        for i, line in enumerate(f):
            line = line.strip()
            if not line:
                continue
            parsed = json.loads(line)
            text = str(parsed.get("text") or "").strip()
            if not text:
                continue
            rows.append(
                {
                    "id": str(parsed.get("id") or f"row_{i}"),
                    "text": text,
                    "metadata": dict(parsed.get("metadata") or {}),
                }
            )
    return rows


def _xlsx_to_chunks(path: Path) -> list[FlatChunk]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl required for xlsx survey parse — pip install openpyxl"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    chunks: list[FlatChunk] = []
    for sheet in wb.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            continue
        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(header)]
        for excel_row, values in enumerate(rows_iter, start=2):
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            parts: list[str] = []
            meta: dict[str, Any] = {
                "excel_row": excel_row,
                "excel_sheet": sheet.title,
                "chunk_type": "excel_row",
                "source_file": path.name,
            }
            persona = None
            for key, val in zip(headers, values, strict=False):
                if val is None or str(val).strip() == "":
                    continue
                text_val = str(val).strip()
                parts.append(f"{key}: {text_val}")
                key_l = key.lower()
                if key_l in {"persona", "persona_name"} and not persona:
                    persona = text_val
                    meta["persona"] = text_val
                if key_l in {"persona_id", "personaid"}:
                    meta["persona_id"] = text_val
            if not parts:
                continue
            body = "\n".join(parts)
            prefix = f"Excel row {excel_row} | sheet {sheet.title}"
            if persona:
                prefix += f" | {persona}"
            chunks.append(
                {
                    "id": f"{sheet.title}_row_{excel_row}",
                    "text": f"{prefix}\n{body}",
                    "metadata": meta,
                }
            )
    wb.close()
    if not chunks:
        raise RuntimeError(f"no data rows found in {path}")
    return chunks


def _docx_to_chunks(path: Path) -> list[FlatChunk]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx required for docx survey parse — pip install python-docx"
        ) from exc

    doc = Document(str(path))
    chunks: list[FlatChunk] = []
    for i, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if not text:
            continue
        chunks.append(
            {
                "id": f"docx_p_{i}",
                "text": text,
                "metadata": {
                    "chunk_type": "docx_paragraph",
                    "paragraph_index": i,
                    "source_file": path.name,
                },
            }
        )
    if not chunks:
        raise RuntimeError(f"no paragraphs found in {path}")
    return chunks


def _write_jsonl_gz(path: Path, chunks: list[FlatChunk]) -> None:
    with gzip.open(path, "wt", encoding="utf-8") as f:
        for chunk in chunks:
            f.write(json.dumps(chunk, ensure_ascii=False))
            f.write("\n")
