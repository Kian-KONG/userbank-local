from __future__ import annotations

import gzip
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from statistics import median
from typing import Any

from .config import get_settings

SheetKind = str

QUESTION_HINTS = re.compile(r"(question|frage|item|statement|item.?text)", re.I)
Q_CODE = re.compile(r"^q\d", re.I)
PERSONA_KEYS = {"persona", "persona_name"}
PERSONA_ID_KEYS = {"persona_id", "personaid"}
CLUSTER_IDS = {"1", "2", "3", "4"}
PLACEHOLDER_CELL = re.compile(r"^\$\{.+\}$")
HEATMAP_SHEET = re.compile(r"heat\s*map", re.I)
MEASURE_CLUSTER = re.compile(r"^(?P<country>.+?)\s+cluster\s+(?P<cluster>[1-4])$", re.I)


@dataclass
class TableBlock:
    sheet: str
    headers: list[str]
    rows: list[tuple[int, list[Any]]]
    start_row: int
    end_row: int
    kind: SheetKind = "data_table"
    column_map: dict[str, dict[str, str]] = field(default_factory=dict)


def parse_tabular_input(
    input_path: Path,
    work: Path,
    track: str = "knowledge",
) -> Path:
    work.mkdir(parents=True, exist_ok=True)
    name = input_path.name.lower()
    long_rows: list[dict[str, Any]] = []
    catalog: list[dict[str, Any]] = []
    if name.endswith(".jsonl.gz") or name.endswith(".jsonl"):
        chunks = _load_jsonl(input_path)
    elif input_path.suffix.lower() in {".xlsx", ".xlsm"}:
        chunks, long_rows, catalog = parse_excel_workbook(input_path, track=track)
    else:
        raise RuntimeError(f"unsupported tabular input: {input_path.name}")
    if not chunks:
        raise RuntimeError(f"no chunks produced from {input_path.name}")
    corpus = {
        "filename": input_path.name,
        "source_type": "excel",
        "chunks": chunks,
    }
    dest = work / "doc_corpus.json"
    dest.write_text(json.dumps(corpus, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if long_rows:
        _write_jsonl_gz(work / "crosstab_long.jsonl.gz", long_rows)
        (work / "question_catalog.json").write_text(
            json.dumps(catalog, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    return dest


def excel_to_chunks(path: Path, track: str = "knowledge") -> list[dict[str, Any]]:
    chunks, _long_rows, _catalog = parse_excel_workbook(path, track=track)
    return chunks


def parse_excel_workbook(
    path: Path, track: str = "knowledge"
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    tables = load_workbook_tables(path)
    chunks: list[dict[str, Any]] = []
    long_rows: list[dict[str, Any]] = []
    for table in tables:
        if _skip_duplicate_crosstab_sheet(table.sheet):
            continue
        table.kind = classify_table(table.headers, table.rows)
        table.column_map = maybe_column_map(table)
        if table.kind == "crosstab":
            records = _crosstab_records(table)
            chunks.extend(_crosstab_chunks_from_records(table, records, path.name))
            long_rows.extend(_unpivot_crosstab_records(records, table.sheet, path.name))
        else:
            chunks.extend(table_to_chunks(table, path.name, track=track))
    catalog = build_question_catalog(long_rows)
    chunks.extend(catalog_chunks(catalog, path.name))
    return chunks, long_rows, catalog


def load_workbook_tables(path: Path) -> list[TableBlock]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl required for xlsx parse — pip install openpyxl") from exc

    wb = load_workbook(path, data_only=True)
    tables: list[TableBlock] = []
    try:
        for sheet in wb.worksheets:
            grid = _sheet_grid(sheet)
            tables.extend(_tables_from_grid(sheet.title, grid))
    finally:
        wb.close()
    if not tables:
        raise RuntimeError(f"no data tables found in {path}")
    return tables


def classify_table(headers: list[str], rows: list[tuple[int, list[Any]]]) -> SheetKind:
    if not headers and not rows:
        return "notes"
    filled_max = 0
    for _excel_row, values in rows:
        filled_max = max(filled_max, sum(1 for v in values if _cell_str(v)))
    if filled_max <= 1 and len(headers) <= 1:
        return "notes"

    if _is_crosstab_headers(headers):
        return "crosstab"
    if _looks_like_codebook(headers):
        return "codebook"
    if _looks_like_kv_form(headers, rows):
        return "kv_form"

    first_header = headers[0] if headers else ""
    first_col = [_cell_str(values[0] if values else "") for _, values in rows]
    other_vals = [
        _cell_str(v)
        for _, values in rows
        for v in values[1:]
        if _cell_str(v)
    ]
    long_questions = bool(first_col) and median(len(v) for v in first_col) >= 40
    low_cardinality = bool(other_vals) and len(set(other_vals)) / max(len(other_vals), 1) <= 0.25
    if QUESTION_HINTS.search(first_header) or (long_questions and low_cardinality):
        return "matrix_questionnaire"
    if any(Q_CODE.match(h.replace(" ", "")) for h in headers):
        return "data_table"
    if not rows:
        return "notes"
    return "data_table"


def maybe_column_map(table: TableBlock) -> dict[str, dict[str, str]]:
    s = get_settings()
    if (
        not s.excel_llm_columns
        or not s.llm_api_key.strip()
        or table.kind not in {"data_table", "matrix_questionnaire"}
    ):
        return {}
    samples = table.rows[:3]
    payload = {
        "sheet": table.sheet,
        "kind": table.kind,
        "headers": table.headers,
        "samples": [
            {table.headers[i]: _cell_str(row[i] if i < len(row) else "") for i in range(len(table.headers))}
            for _, row in samples
        ],
    }
    try:
        return _llm_column_map(payload)
    except Exception as exc:  # noqa: BLE001
        print(f"==> excel column map skipped: {exc}")
        return {}


def table_to_chunks(
    table: TableBlock,
    source_file: str,
    track: str = "knowledge",
) -> list[dict[str, Any]]:
    headers = [_mapped_header(table, h) for h in table.headers]
    if table.kind == "notes":
        body = "\n".join(
            " | ".join(_cell_str(v) for v in row if _cell_str(v))
            for _, row in table.rows
        ).strip()
        if not body:
            return []
        return [_chunk(
            f"{table.sheet}_notes",
            f"Sheet {table.sheet}\n{body}",
            {
                "excel_sheet": table.sheet,
                "sheet_kind": "notes",
                "chunk_type": "excel_notes",
                "source_file": source_file,
            },
        )]

    if table.kind == "codebook":
        return _codebook_chunks(table, headers, source_file)
    if table.kind == "crosstab":
        return _crosstab_chunks(table, source_file)

    chunks: list[dict[str, Any]] = []
    for excel_row, values in table.rows:
        parts: list[str] = []
        meta: dict[str, Any] = {
            "excel_row": excel_row,
            "excel_sheet": table.sheet,
            "sheet_kind": table.kind,
            "chunk_type": "excel_row" if table.kind != "matrix_questionnaire" else "excel_question",
            "source_file": source_file,
        }
        for key, val in zip(headers, values, strict=False):
            text_val = _cell_str(val)
            if not text_val:
                continue
            parts.append(f"{key}: {text_val}")
            if track == "survey":
                key_l = key.lower()
                if key_l in PERSONA_KEYS and "persona" not in meta:
                    meta["persona"] = text_val
                if key_l in PERSONA_ID_KEYS:
                    meta["persona_id"] = text_val
        if not parts:
            continue
        prefix = f"Excel row {excel_row} | sheet {table.sheet} | {table.kind}"
        if meta.get("persona"):
            prefix += f" | {meta['persona']}"
        chunks.append(_chunk(f"{table.sheet}_row_{excel_row}", f"{prefix}\n" + "\n".join(parts), meta))
    return chunks


def _codebook_chunks(
    table: TableBlock,
    headers: list[str],
    source_file: str,
) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    current_var = ""
    parts: list[str] = []
    start_row = 0
    for excel_row, values in table.rows:
        row_map = {
            headers[i]: _cell_str(values[i] if i < len(values) else "")
            for i in range(len(headers))
        }
        var = row_map.get(headers[0], "") if headers else ""
        if var and var != current_var and parts:
            chunks.append(_chunk(
                f"{table.sheet}_var_{start_row}",
                f"Sheet {table.sheet} | codebook\n" + "\n".join(parts),
                {
                    "excel_sheet": table.sheet,
                    "excel_row": start_row,
                    "sheet_kind": "codebook",
                    "chunk_type": "excel_codebook",
                    "source_file": source_file,
                    "variable": current_var,
                },
            ))
            parts = []
        if var:
            current_var = var
            start_row = excel_row
        line = " | ".join(f"{k}: {v}" for k, v in row_map.items() if v)
        if line:
            parts.append(line)
    if parts:
        chunks.append(_chunk(
            f"{table.sheet}_var_{start_row}",
            f"Sheet {table.sheet} | codebook\n" + "\n".join(parts),
            {
                "excel_sheet": table.sheet,
                "excel_row": start_row,
                "sheet_kind": "codebook",
                "chunk_type": "excel_codebook",
                "source_file": source_file,
                "variable": current_var,
            },
        ))
    return chunks


def _mapped_header(table: TableBlock, header: str) -> str:
    mapped = table.column_map.get(header) or {}
    label = str(mapped.get("label") or "").strip()
    return label or header


def _chunk(chunk_id: str, text: str, metadata: dict[str, Any]) -> dict[str, Any]:
    return {"id": chunk_id, "text": text, "metadata": metadata}


def _sheet_grid(sheet: Any) -> list[list[Any]]:
    max_row = sheet.max_row or 0
    max_col = sheet.max_column or 0
    if max_row < 1 or max_col < 1:
        return []
    grid = [[None] * max_col for _ in range(max_row)]
    for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            grid[cell.row - 1][cell.column - 1] = cell.value
    for merged in sheet.merged_cells.ranges:
        value = grid[merged.min_row - 1][merged.min_col - 1]
        for row_i in range(merged.min_row, merged.max_row + 1):
            for col_i in range(merged.min_col, merged.max_col + 1):
                grid[row_i - 1][col_i - 1] = value
    return grid


def _tables_from_grid(sheet_name: str, grid: list[list[Any]]) -> list[TableBlock]:
    if not grid:
        return []
    tables: list[TableBlock] = []
    i = 0
    while i < len(grid):
        while i < len(grid) and _is_empty_row(grid[i]):
            i += 1
        if i >= len(grid):
            break
        header_start = _skip_title_rows(grid, i)
        header_end = _header_band_end(grid, header_start)
        headers = _join_headers(grid[header_start:header_end])
        data_start = header_end
        data_end = data_start
        while data_end < len(grid) and not _is_empty_row(grid[data_end]):
            data_end += 1
        rows = [
            (idx + 1, list(grid[idx]))
            for idx in range(data_start, data_end)
            if not _is_empty_row(grid[idx])
        ]
        if headers or rows:
            tables.append(
                TableBlock(
                    sheet=sheet_name,
                    headers=headers or [f"col_{n}" for n in range(_row_width(rows))],
                    rows=_trim_rows(rows, max(len(headers), 1)),
                    start_row=header_start + 1,
                    end_row=data_end,
                )
            )
        i = data_end + 1
    return tables


def _skip_title_rows(grid: list[list[Any]], start: int) -> int:
    i = start
    while i < len(grid) - 1:
        filled = _filled_count(grid[i])
        nxt = _filled_count(grid[i + 1])
        if filled <= 1 and nxt >= max(2, filled + 1) and not _mostly_numeric(grid[i + 1]):
            i += 1
            continue
        break
    return i


def _header_band_end(grid: list[list[Any]], start: int) -> int:
    if start + 1 < len(grid) and _is_cluster_id_row(grid[start + 1]):
        return start + 2
    end = start + 1
    while end < len(grid) and _looks_like_header(grid[end]) and not _is_empty_row(grid[end]):
        if _mostly_numeric(grid[end]):
            break
        end += 1
        if end - start >= 4:
            break
    return max(end, start + 1)


def _join_headers(header_rows: list[list[Any]]) -> list[str]:
    if not header_rows:
        return []
    if len(header_rows) >= 2 and _is_cluster_id_row(header_rows[1]):
        return _crosstab_headers(header_rows[0], header_rows[1])
    width = max(len(row) for row in header_rows)
    headers: list[str] = []
    for col in range(width):
        parts: list[str] = []
        for row in header_rows:
            value = _cell_str(row[col] if col < len(row) else "")
            if value and value not in parts:
                parts.append(value)
        headers.append(" / ".join(parts) if parts else f"col_{col}")
    while len(headers) > 1 and headers[-1].startswith("col_"):
        headers.pop()
    return headers


def _skip_duplicate_crosstab_sheet(sheet: str) -> bool:
    return bool(HEATMAP_SHEET.search(sheet or ""))


def _is_cluster_id_row(row: list[Any]) -> bool:
    filled = [_cell_str(v) for v in row if _cell_str(v)]
    if len(filled) < 8:
        return False
    ids: list[str] = []
    for value in filled:
        if not _is_number(value):
            return False
        number = float(value.replace(",", ""))
        if number not in {1.0, 2.0, 3.0, 4.0}:
            return False
        ids.append(str(int(number)))
    return ids[:4] == ["1", "2", "3", "4"] and ids[4:8] == ["1", "2", "3", "4"]


def _is_crosstab_headers(headers: list[str]) -> bool:
    cluster_headers = [h for h in headers if " cluster " in h]
    return len(cluster_headers) >= 8


def _crosstab_stub_names(count: int) -> list[str]:
    names = ["question", "option", "metric", "Overall"]
    if count <= len(names):
        return names[:count]
    extra = [f"col_{index}" for index in range(len(names), count)]
    return names + extra


def _crosstab_headers(country_row: list[Any], cluster_row: list[Any]) -> list[str]:
    width = max(len(country_row), len(cluster_row))
    countries = [_cell_str(country_row[col] if col < len(country_row) else "") for col in range(width)]
    clusters = [_cell_str(cluster_row[col] if col < len(cluster_row) else "") for col in range(width)]
    first_measure = next(
        (
            col
            for col in range(width)
            if countries[col] or _cluster_token(clusters[col])
        ),
        width,
    )
    stub_names = _crosstab_stub_names(first_measure)
    headers: list[str] = []
    last_country = ""
    last_filled = 0
    for col in range(width):
        if col < first_measure:
            headers.append(stub_names[col])
            last_filled = len(headers)
            continue
        if countries[col]:
            last_country = countries[col]
        cluster = _cluster_token(clusters[col])
        if cluster:
            label = f"{last_country} cluster {cluster}" if last_country else f"cluster {cluster}"
            headers.append(label)
            last_filled = len(headers)
        elif countries[col]:
            headers.append(countries[col])
            last_filled = len(headers)
        else:
            headers.append(f"col_{col}")
    return headers[:last_filled]


def _cluster_token(value: str) -> str:
    if not value or not _is_number(value):
        return ""
    number = float(value.replace(",", ""))
    if number not in {1.0, 2.0, 3.0, 4.0}:
        return ""
    return str(int(number))


def _crosstab_chunks(table: TableBlock, source_file: str) -> list[dict[str, Any]]:
    return _crosstab_chunks_from_records(table, _crosstab_records(table), source_file)


def _crosstab_records(table: TableBlock) -> list[dict[str, Any]]:
    headers = table.headers
    question_i = headers.index("question") if "question" in headers else None
    option_i = headers.index("option") if "option" in headers else None
    metric_i = headers.index("metric") if "metric" in headers else None
    measure_indexes = [
        index
        for index, header in enumerate(headers)
        if index not in {question_i, option_i, metric_i}
    ]
    records: list[dict[str, Any]] = []
    question = ""
    option = ""
    for excel_row, values in table.rows:
        padded = list(values) + [None] * max(0, len(headers) - len(values))
        row_question = _cell_str(padded[question_i]) if question_i is not None else ""
        row_option = _cell_str(padded[option_i]) if option_i is not None else ""
        metric = _cell_str(padded[metric_i]) if metric_i is not None else ""
        measures = _crosstab_measures(headers, padded, measure_indexes)
        if option_i is None:
            if row_question and not measures:
                question = row_question
                continue
            if row_question:
                option = row_question
        else:
            if row_question:
                question = row_question
            if row_option:
                option = row_option
        if _skip_crosstab_record(question, option or row_question):
            continue
        if not measures or _measures_all_zero(measures):
            continue
        if not option:
            option = "Base n"
        records.append(
            {
                "excel_row": excel_row,
                "question": question,
                "option": option,
                "metric": metric,
                "measures": measures,
            }
        )
    return records


def _crosstab_chunks_from_records(
    table: TableBlock,
    records: list[dict[str, Any]],
    source_file: str,
) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    index = 0
    while index < len(records):
        current = records[index]
        paired = None
        if index + 1 < len(records):
            nxt = records[index + 1]
            if (
                _metric_kind(current["metric"]) == "value"
                and _metric_kind(nxt["metric"]) == "percent"
                and nxt["question"] == current["question"]
                and nxt["option"] == current["option"]
            ):
                paired = nxt
        text = _format_crosstab_chunk(table.sheet, current, paired)
        if text:
            meta = {
                "excel_row": current["excel_row"],
                "excel_sheet": table.sheet,
                "sheet_kind": "crosstab",
                "chunk_type": "excel_crosstab",
                "source_file": source_file,
                "question": current["question"],
                "option": current["option"],
            }
            chunks.append(
                _chunk(
                    f"{table.sheet}_q_{current['excel_row']}",
                    text,
                    meta,
                )
            )
        index += 2 if paired is not None else 1
    return chunks


def _unpivot_crosstab_records(
    records: list[dict[str, Any]],
    sheet: str,
    source_file: str,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for record in records:
        metric = "pct" if _metric_kind(str(record.get("metric") or "")) == "percent" else "count"
        for header, raw in record.get("measures") or []:
            if not _is_number(raw):
                continue
            country, cluster = _split_measure_header(header)
            if not country:
                continue
            rows.append(
                {
                    "source_file": source_file,
                    "sheet": sheet,
                    "question": record["question"],
                    "option": record["option"],
                    "country": country,
                    "cluster": cluster,
                    "metric": metric,
                    "value": float(raw.replace(",", "")),
                }
            )
    return rows


def _split_measure_header(header: str) -> tuple[str, int | None]:
    match = MEASURE_CLUSTER.match((header or "").strip())
    if match:
        return match.group("country").strip(), int(match.group("cluster"))
    return (header or "").strip(), None


def build_question_catalog(long_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for row in long_rows:
        question = str(row.get("question") or "").strip()
        sheet = str(row.get("sheet") or "").strip()
        if not question:
            continue
        entry = grouped.setdefault(
            (sheet, question),
            {
                "sheet": sheet,
                "question": question,
                "options": set(),
                "countries": set(),
                "clusters": set(),
            },
        )
        option = str(row.get("option") or "").strip()
        country = str(row.get("country") or "").strip()
        cluster = row.get("cluster")
        if option:
            entry["options"].add(option)
        if country:
            entry["countries"].add(country)
        if cluster is not None and str(cluster).strip() != "":
            entry["clusters"].add(int(cluster))
    catalog: list[dict[str, Any]] = []
    for entry in grouped.values():
        options = sorted(entry["options"])
        countries = sorted(entry["countries"])
        clusters = sorted(entry["clusters"])
        catalog.append(
            {
                "sheet": entry["sheet"],
                "question": entry["question"],
                "options": options,
                "countries": countries,
                "clusters": clusters,
                "option_count": len(options),
            }
        )
    return catalog


def catalog_chunks(catalog: list[dict[str, Any]], source_file: str) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    for index, entry in enumerate(catalog):
        question = entry["question"]
        options = entry.get("options") or []
        countries = entry.get("countries") or []
        preview = "; ".join(str(item) for item in options[:12])
        if len(options) > 12:
            preview += "; …"
        text = (
            f"Survey catalog | {source_file}\n"
            f"Question: {question}\n"
            f"Sheet: {entry.get('sheet')}\n"
            f"Options ({entry.get('option_count') or len(options)}): {preview}\n"
            f"Countries: {', '.join(str(item) for item in countries)}\n"
            "Look up counts (metric=count) and column percentages (metric=pct) "
            "by country and cluster 1-4 in the survey table."
        )
        chunks.append(
            _chunk(
                f"catalog_{index}_{entry.get('sheet')}",
                text,
                {
                    "excel_sheet": entry.get("sheet"),
                    "sheet_kind": "crosstab",
                    "chunk_type": "excel_catalog",
                    "source_file": source_file,
                    "question": question,
                },
            )
        )
    return chunks


def _skip_crosstab_record(question: str, option: str) -> bool:
    blobs = (question, option)
    if any(PLACEHOLDER_CELL.match(text) for text in blobs):
        return True
    if question.strip().lower() == "hidden for country":
        return True
    return False


def _crosstab_measures(
    headers: list[str],
    values: list[Any],
    measure_indexes: list[int],
) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    for index in measure_indexes:
        text = _cell_str(values[index] if index < len(values) else "")
        if not text:
            continue
        out.append((headers[index], text))
    return out


def _measures_all_zero(measures: list[tuple[str, str]]) -> bool:
    if not measures:
        return True
    for _header, raw in measures:
        if not _is_number(raw):
            return False
        if float(raw.replace(",", "")) != 0:
            return False
    return True


def _metric_kind(metric: str) -> str:
    lowered = metric.strip().lower()
    if lowered == "value":
        return "value"
    if "percent" in lowered:
        return "percent"
    return lowered


def _format_crosstab_chunk(
    sheet: str,
    record: dict[str, Any],
    percent_record: dict[str, Any] | None,
) -> str:
    lines = [f"Crosstab | {sheet}"]
    question = str(record.get("question") or "").strip()
    option = str(record.get("option") or "").strip()
    if question:
        lines.append(f"Question: {question}")
    if option:
        lines.append(f"Answer: {option}")
    value_lines = _measure_lines(record.get("measures") or [], as_percent=False)
    if value_lines:
        heading = "Value:" if percent_record is not None or _metric_kind(record.get("metric") or "") == "value" else "Counts:"
        if _metric_kind(record.get("metric") or "") == "percent" and percent_record is None:
            heading = "Column percentage:"
            value_lines = _measure_lines(record.get("measures") or [], as_percent=True)
        lines.append(heading)
        lines.extend(value_lines)
    if percent_record is not None:
        pct_lines = _measure_lines(percent_record.get("measures") or [], as_percent=True)
        if pct_lines:
            lines.append("Column percentage:")
            lines.extend(pct_lines)
    if len(lines) <= 3:
        return ""
    return "\n".join(lines)


def _measure_lines(measures: list[tuple[str, str]], *, as_percent: bool) -> list[str]:
    lines: list[str] = []
    for header, raw in measures:
        display = _format_measure(raw, as_percent=as_percent)
        if not display:
            continue
        lines.append(f"- {header}: {display}")
    return lines


def _format_measure(raw: str, *, as_percent: bool) -> str:
    if not raw:
        return ""
    if not as_percent or not _is_number(raw):
        return raw
    number = float(raw.replace(",", ""))
    if 0 <= number <= 1:
        return f"{number * 100:.1f}%"
    return raw


def _trim_rows(rows: list[tuple[int, list[Any]]], width: int) -> list[tuple[int, list[Any]]]:
    out: list[tuple[int, list[Any]]] = []
    for excel_row, values in rows:
        padded = list(values) + [None] * max(0, width - len(values))
        out.append((excel_row, padded[:width]))
    return out


def _row_width(rows: list[tuple[int, list[Any]]]) -> int:
    if not rows:
        return 1
    return max(len(values) for _, values in rows)


def _looks_like_codebook(headers: list[str]) -> bool:
    if len(headers) < 2:
        return False
    blob = " ".join(headers)
    hits = 0
    if re.search(r"(variable|var_?name)\b", blob, re.I):
        hits += 1
    if re.search(r"\bcode\b", blob, re.I):
        hits += 1
    if re.search(
        r"(value.?label|\blabel\b|meaning|beschreibung|bezeichnung)", blob, re.I
    ):
        hits += 1
    return hits >= 2


def _looks_like_kv_form(
    headers: list[str],
    rows: list[tuple[int, list[Any]]],
) -> bool:
    if len(headers) != 2 or not rows:
        return False
    left = headers[0]
    right = headers[1]
    if re.search(r"(key|field|label|attribute|item)", left, re.I) and re.search(
        r"(value|val|content|answer)", right, re.I
    ):
        return True
    first_col = [_cell_str(values[0] if values else "") for _, values in rows]
    filled = [v for v in first_col if v]
    if len(filled) < 3 or len(set(filled)) != len(filled):
        return False
    if any(Q_CODE.match(h.replace(" ", "")) for h in headers):
        return False
    return median(len(v) for v in filled) <= 24


def _looks_like_header(row: list[Any]) -> bool:
    filled = [_cell_str(v) for v in row if _cell_str(v)]
    if len(filled) < 2:
        return False
    numeric = sum(1 for v in filled if _is_number(v))
    return numeric / len(filled) <= 0.3


def _mostly_numeric(row: list[Any]) -> bool:
    filled = [_cell_str(v) for v in row if _cell_str(v)]
    if not filled:
        return False
    numeric = sum(1 for v in filled if _is_number(v))
    return numeric / len(filled) >= 0.5


def _filled_count(row: list[Any]) -> int:
    return sum(1 for v in row if _cell_str(v))


def _is_empty_row(row: list[Any]) -> bool:
    return _filled_count(row) == 0


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_number(value: str) -> bool:
    try:
        float(value.replace(",", ""))
        return True
    except ValueError:
        return False


def _write_jsonl_gz(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(path, "wt", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def _load_jsonl(path: Path) -> list[dict[str, Any]]:
    opener = gzip.open if path.name.lower().endswith(".gz") else open
    rows: list[dict[str, Any]] = []
    with opener(path, "rt", encoding="utf-8") as handle:
        for index, line in enumerate(handle):
            line = line.strip()
            if not line:
                continue
            parsed = json.loads(line)
            text = str(parsed.get("text") or "").strip()
            if not text:
                continue
            rows.append(
                {
                    "id": str(parsed.get("id") or f"row_{index}"),
                    "text": text,
                    "metadata": dict(parsed.get("metadata") or {}),
                }
            )
    return rows


def _llm_column_map(payload: dict[str, Any]) -> dict[str, dict[str, str]]:
    import httpx

    s = get_settings()
    prompt = (
        "You label Excel columns for retrieval. Return JSON only: "
        '{"columns": {"<header>": {"label": "human name", "role": "id|measure|category|question|other", "unit": ""}}}.\n'
        "Keep labels short. Do not invent columns.\n\n"
        + json.dumps(payload, ensure_ascii=False)
    )
    with httpx.Client(timeout=60.0) as client:
        resp = client.post(
            s.llm_api_url,
            json={
                "model": s.llm_synth_model,
                "messages": [{"role": "user", "content": prompt}],
                "max_tokens": 2048,
            },
            headers={
                "Authorization": f"Bearer {s.llm_api_key}",
                "Content-Type": "application/json",
            },
        )
    if resp.is_error:
        raise RuntimeError(f"Excel map HTTP {resp.status_code}: {resp.text[:300]}")
    raw = resp.json()["choices"][0]["message"]["content"]
    start = raw.find("{")
    end = raw.rfind("}")
    if start < 0 or end <= start:
        return {}
    parsed = json.loads(raw[start : end + 1])
    columns = parsed.get("columns") if isinstance(parsed, dict) else None
    if not isinstance(columns, dict):
        return {}
    out: dict[str, dict[str, str]] = {}
    for key, value in columns.items():
        if isinstance(value, dict):
            out[str(key)] = {
                "label": str(value.get("label") or key),
                "role": str(value.get("role") or "other"),
                "unit": str(value.get("unit") or ""),
            }
    return out
